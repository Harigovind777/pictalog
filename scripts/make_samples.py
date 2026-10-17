"""Build the sample catalogs in samples/: one of each kind of file a company might upload.

uv run python scripts/make_samples.py
"""

import io
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT = Path(__file__).resolve().parent.parent / "samples"


def apparel_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Spring 2026"
    ws.append(["Kerala Weaves Pvt Ltd — Spring 2026 price list"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Item Code", "Product Name", "Category", "Colour", "Sizes", "Fabric", "MRP (₹)", "Remarks"])
    ws.append(["KW-101", "Classic Crew Neck T-Shirt", "T-Shirts", "Red/Navy", "S, M, L", "Cotton", 499, "Bestseller"])
    ws.append(["KW-205", "Hooded Sweatshirt", "Sweatshirts", "Heather Grey", "M, L", "Cotton fleece", 1299.0, ""])
    ws.append(["KW-310", "Kasavu Mundu", "Ethnic Wear", "Off-white with gold border", "Free Size", "Handloom cotton", 899, "Handwoven"])
    ws.append(["KW-999", "KW-999", "", "", "", "", 250, "Sample row with only a code"])
    wb.save(OUT / "apparel.xlsx")


def kitchen_csv() -> None:
    rows = [
        "Product;Price;Size;Colour;Material",
        "Mug en céramique;4,50 €;350 ml;Blanc;Céramique",
        "Bouteille isotherme;19,90 €;1 L;Noir mat;Acier inoxydable",
        "Planche à découper;24,00 €;40 x 30 cm;Naturel;Bois d'acacia",
        "Théière;32,00 €;1,2 L;Bleu;Fonte",
    ]
    (OUT / "kitchen.csv").write_bytes(("\n".join(rows) + "\n").encode("cp1252"))  # what French Excel saves


def furniture_docx() -> None:
    doc = Document()
    doc.add_heading("Teakwood Living — new arrivals", level=1)
    doc.add_paragraph(
        "Our Nilambur teak dining table seats six and is finished in a warm natural oil. It costs ₹48,000. "
        "The matching dining chair has a woven cane seat and sells for ₹6,500 each."
    )
    doc.add_paragraph(
        "New this season: a three-seater sofa in mustard yellow linen with solid teak legs, ₹62,000, "
        "and a round brass-top coffee table, 80 cm across, for ₹18,500."
    )
    doc.add_paragraph("Delivery in Kochi and Thrissur is free. Call 0484 000 0000 for orders.")
    doc.save(OUT / "furniture.docx")


def electronics_pdf() -> None:
    buf = io.BytesIO()
    styles = getSampleStyleSheet()
    table = Table(
        [
            ["Model", "Description", "Colour", "Price"],
            ["Wireless Earbuds Pro", "Noise-cancelling earbuds with charging case", "White", "₹4,999"],
            ["Smart Watch 2", "Fitness watch with AMOLED screen", "Midnight Black", "₹7,499"],
            ["Bluetooth Speaker Mini", "Water-resistant portable speaker", "Teal", "₹2,299"],
        ],
        repeatRows=1,
    )
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]))
    SimpleDocTemplate(buf, pagesize=A4).build([Paragraph("Volt Gadgets — dealer price list", styles["Title"]), Spacer(1, 12), table])
    (OUT / "electronics.pdf").write_bytes(buf.getvalue())


def pricelist_jpg() -> None:
    """A phone photo of a printed price list: slightly rotated paper on a table."""
    paper = Image.new("RGB", (900, 700), "#fbfaf6")
    draw = ImageDraw.Draw(paper)
    try:
        title = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 40)
        body = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 30)
    except OSError:
        title = body = ImageFont.load_default()
    draw.text((60, 50), "Spice Garden — Price List", fill="#222", font=title)
    lines = [
        ("Black Pepper (Tellicherry), 250 g", "Rs. 320"),
        ("Green Cardamom, 100 g", "Rs. 450"),
        ("Cinnamon Sticks, 200 g", "Rs. 180"),
        ("Turmeric Powder, 500 g", "Rs. 140"),
    ]
    for i, (name, price) in enumerate(lines):
        y = 160 + i * 90
        draw.text((60, y), name, fill="#333", font=body)
        draw.text((690, y), price, fill="#333", font=body)
        draw.line((60, y + 55, 840, y + 55), fill="#ddd", width=2)
    photo = Image.new("RGB", (1100, 900), "#8a6f52")
    photo.paste(paper.rotate(2.5, expand=True, fillcolor="#8a6f52"), (90, 80))
    photo.filter(ImageFilter.GaussianBlur(0.6)).save(OUT / "pricelist.jpg", quality=88)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    apparel_xlsx()
    kitchen_csv()
    furniture_docx()
    electronics_pdf()
    pricelist_jpg()
    print("\n".join(sorted(p.name for p in OUT.iterdir())))
