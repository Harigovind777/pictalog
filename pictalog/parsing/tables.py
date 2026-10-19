import csv
import io
import json
import re
import unicodedata

import pandas as pd
import pdfplumber
from docx import Document

from pictalog.parsing.cells import clean_cell, decode_text
from pictalog.parsing.detect import FileKind

Table = list[list[str]]

DELIMITERS = [",", ";", "\t", "|"]


def read_tables(kind: FileKind, data: bytes) -> list[Table]:
    """Every table found in a tidy file, as rows of cleaned text. The header row is not identified here."""
    match kind:
        case "spreadsheet":
            return _spreadsheet(data)
        case "csv":
            return [_csv(decode_text(data))]
        case "json":
            return _json(data)
        case "docx":
            return _docx(data)
        case "pdf":
            return [] if _broken_text_layer(_pdf_text(data)) else _pdf(data)
        case "text":
            text = decode_text(data)
            return [_csv(text)] if _tab_separated(text) else []
    return []


def _tab_separated(text: str) -> bool:
    """Spreadsheets saved as text ("Unicode Text" in Excel) are tab-separated; prose is left for the AI reader."""
    lines = [line for line in text.splitlines() if line.strip()]
    return len(lines) >= 2 and "\t" in lines[0] and sum("\t" in line for line in lines) >= len(lines) * 0.8


def extract_text(kind: FileKind, data: bytes) -> str:
    """Plain text for the AI reader; "" when there is none (photos, scanned PDFs)."""
    match kind:
        case "docx":
            doc = Document(io.BytesIO(data))
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                lines += [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            return "\n".join(lines)
        case "pdf":
            text = _pdf_text(data)
            return "" if _broken_text_layer(text) else text
        case "text" | "csv":
            return decode_text(data)
        case "json":
            return decode_text(data)
    return ""


def _pdf_text(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages).strip()


_INDIAN_LETTER = re.compile(r"[\u0900-\u0dff\u1c50-\u1c7f\uabc0-\uabff]")
_LOST_LETTER = re.compile(r"[\x00\ufffd\ue000-\uf8ff]")  # nothing, "unknown", or the font's private glyph numbers
_ARABIC_AS_DRAWN = re.compile(r"[\ufb50-\ufdff\ufe70-\ufefc]")  # Urdu letters in their joined shapes, right to left


def _broken_text_layer(text: str) -> bool:
    """Indian-script text in many PDFs doesn't survive extraction: letters drawn joined come out as nothing or as
    private characters, and vowel signs in the order they're drawn (புடைவ for புடவை). Those PDFs are read as pictures,
    by OCR, instead."""
    if _ARABIC_AS_DRAWN.search(text):
        return True
    if not _INDIAN_LETTER.search(text):
        return False
    return bool(_LOST_LETTER.search(text)) or any(unicodedata.category(word[0])[0] == "M" for word in text.split())


def _spreadsheet(data: bytes) -> list[Table]:
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=object)
    return [
        [[clean_cell(v) for v in row] for row in frame.itertuples(index=False)]
        for frame in sheets.values()
        if not frame.empty
    ]


def _csv(text: str) -> Table:
    first_line = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = max(DELIMITERS, key=first_line.count) if first_line else ","
    if first_line.count(delimiter) == 0:
        delimiter = ","
    return [[clean_cell(v) for v in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _json(data: bytes) -> list[Table]:
    try:
        value = json.loads(decode_text(data))
    except json.JSONDecodeError:
        return []
    records = _find_records(value)
    if not records:
        return []
    header: list[str] = []
    for record in records:
        header += [k for k in record if k not in header]
    return [[header] + [[clean_cell(_scalar(r.get(k))) for k in header] for r in records]]


def _find_records(value: object) -> list[dict] | None:
    def is_records(v: object) -> bool:
        return isinstance(v, list) and bool(v) and all(isinstance(x, dict) for x in v)

    if is_records(value):
        return value  # type: ignore[return-value]
    if isinstance(value, dict):
        return next((v for v in value.values() if is_records(v)), None)
    return None


def _scalar(value: object) -> object:
    return json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value


def _docx(data: bytes) -> list[Table]:
    doc = Document(io.BytesIO(data))
    return [[[clean_cell(cell.text) for cell in row.cells] for row in table.rows] for table in doc.tables]


def _pdf(data: bytes) -> list[Table]:
    tables: list[Table] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            for raw in page.extract_tables():
                table = [[clean_cell(v) for v in row] for row in raw]
                if tables and table and table[0] == tables[-1][0]:
                    tables[-1].extend(table[1:])  # the same table continued on the next page
                else:
                    tables.append(table)
    return tables
