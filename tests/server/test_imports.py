import io

from docx import Document
from openpyxl import Workbook
from sqlalchemy import select

from pictalog.models import Job

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def docx_text(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


SHEET = [["Acme Spring 2026"], [], ["Product", "Price", "Size", "Remarks"], ["Tee", 499, "S, M", "Soft cotton"], ["Mug", 199, "350 ml", ""]]


def new_catalog(client, headers) -> int:
    return client.post("/api/catalogs", json={"name": "Spring"}, headers=headers).json()["id"]


def upload(client, headers, cid, name, data, mime="application/octet-stream"):
    return client.post(f"/api/catalogs/{cid}/imports", files={"file": (name, data, mime)}, headers=headers)


def test_upload_xlsx_ready_for_review(client, signup):
    h = signup()
    cid = new_catalog(client, h)
    r = upload(client, h, cid, "spring.xlsx", xlsx(SHEET), XLSX)
    assert r.status_code == 201
    body = r.json()
    assert body["status"] == "ready_for_review" and body["method"] == "table"
    assert body["headers"] == ["Product", "Price", "Size", "Remarks"]
    assert body["mapping"] == {"Product": "name", "Price": "price", "Size": "size", "Remarks": None}
    assert [(i["name"], i["size"], i["price_text"]) for i in body["items"]] == [
        ("Tee", "S", "499"), ("Tee", "M", "499"), ("Mug", "350 ml", "199"),
    ]
    assert body["items"][0]["extra"] == {"Remarks": "Soft cotton"}
    assert body["diff"]["new"] == 3
    assert client.get(f"/api/imports/{body['id']}", headers=h).json()["status"] == "ready_for_review"


def test_upload_free_text_docx_needs_ai(client, signup, db):
    h = signup()
    cid = new_catalog(client, h)
    r = upload(client, h, cid, "range.docx", docx_text("Our oak dining table costs 25,000 and seats six."))
    assert r.status_code == 201
    body = r.json()
    assert body["status"] == "needs_ai" and body["method"] == "ai" and body["diff"] is None
    [job] = db.scalars(select(Job)).all()
    assert job.kind == "extract" and job.status == "queued" and job.import_id == body["id"]
    assert "oak dining table" in job.payload["text"]


def test_upload_photo_needs_ai_without_text(client, signup, db):
    h = signup()
    cid = new_catalog(client, h)
    r = upload(client, h, cid, "pricelist.jpg", b"\xff\xd8\xff fake jpeg", "image/jpeg")
    assert r.status_code == 201 and r.json()["status"] == "needs_ai"
    [job] = db.scalars(select(Job)).all()
    assert job.payload["text"] is None


def test_spreadsheet_without_header_goes_to_ai(client, signup, db):
    h = signup()
    cid = new_catalog(client, h)
    r = upload(client, h, cid, "odd.xlsx", xlsx([["Tee shirt red", "499"], ["Mug white", "199"]]), XLSX)
    assert r.json()["status"] == "needs_ai"
    [job] = db.scalars(select(Job)).all()
    assert "Tee shirt red | 499" in job.payload["text"]


def test_upload_too_large_413(client, signup, app):
    h = signup()
    cid = new_catalog(client, h)
    app.state.settings.max_upload_mb = 1
    r = upload(client, h, cid, "big.csv", b"x" * (1024 * 1024 + 10))
    assert r.status_code == 413
    assert "1 MB" in r.json()["detail"]


def test_upload_doc_415_mentions_docx(client, signup):
    h = signup()
    cid = new_catalog(client, h)
    r = upload(client, h, cid, "old.doc", b"data")
    assert r.status_code == 415 and ".docx" in r.json()["detail"]


def test_corrupt_file_400(client, signup):
    h = signup()
    cid = new_catalog(client, h)
    r = upload(client, h, cid, "broken.xlsx", b"not really excel", XLSX)
    assert r.status_code == 400


def test_empty_document_400(client, signup):
    h = signup()
    cid = new_catalog(client, h)
    r = upload(client, h, cid, "empty.txt", b"   ")
    assert r.status_code == 400


def test_change_mapping_rebuilds_items(client, signup):
    h = signup()
    cid = new_catalog(client, h)
    iid = upload(client, h, cid, "spring.xlsx", xlsx(SHEET), XLSX).json()["id"]
    r = client.put(f"/api/imports/{iid}/mapping", json={"mapping": {"Remarks": "description"}}, headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["mapping"]["Remarks"] == "description" and body["mapping"]["Product"] == "name"
    assert body["items"][0]["description"] == "Soft cotton" and body["items"][0]["extra"] == {}


def test_change_mapping_rejects_field_used_twice(client, signup):
    h = signup()
    cid = new_catalog(client, h)
    iid = upload(client, h, cid, "spring.xlsx", xlsx(SHEET), XLSX).json()["id"]
    r = client.put(f"/api/imports/{iid}/mapping", json={"mapping": {"Remarks": "name"}}, headers=h)
    assert r.status_code == 422


def test_mapping_not_allowed_for_ai_imports(client, signup):
    h = signup()
    cid = new_catalog(client, h)
    iid = upload(client, h, cid, "range.docx", docx_text("Oak table 25,000")).json()["id"]
    assert client.put(f"/api/imports/{iid}/mapping", json={"mapping": {}}, headers=h).status_code == 409


def test_apply_creates_products_and_is_single_use(client, signup):
    h = signup()
    cid = new_catalog(client, h)
    imp = upload(client, h, cid, "spring.xlsx", xlsx(SHEET), XLSX).json()
    items = imp["items"]
    items[2]["name"] = "Coffee Mug"  # edited on the check screen
    r = client.post(f"/api/imports/{imp['id']}/apply", json={"items": items}, headers=h)
    assert r.status_code == 200
    catalog = r.json()
    assert [p["name"] for p in catalog["products"]] == ["Tee", "Coffee Mug"]
    assert catalog["counts"] == {"queued": 3}
    assert client.get(f"/api/imports/{imp['id']}", headers=h).json()["status"] == "applied"
    assert client.post(f"/api/imports/{imp['id']}/apply", json={"items": items}, headers=h).status_code == 409


def test_other_company_import_404(client, signup):
    a = signup("a@x.com")
    b = signup("b@x.com")
    cid = new_catalog(client, a)
    imp = upload(client, a, cid, "spring.xlsx", xlsx(SHEET), XLSX).json()
    assert upload(client, b, cid, "spring.xlsx", xlsx(SHEET), XLSX).status_code == 404
    assert client.get(f"/api/imports/{imp['id']}", headers=b).status_code == 404
    assert client.put(f"/api/imports/{imp['id']}/mapping", json={"mapping": {}}, headers=b).status_code == 404
    assert client.post(f"/api/imports/{imp['id']}/apply", json={"items": []}, headers=b).status_code == 404
